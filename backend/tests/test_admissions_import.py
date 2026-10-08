"""
Automated tests for Government Admission Ingestion Engine.
Tests:
- Upload and staging
- Validation (valid, invalid, duplicate)
- Partial failure resilience
- Atomic commit into core Student and Admission entities
- Role-based authorization
"""
import datetime
import io
import pytest
from rest_framework import status
from rest_framework.test import APIClient

from apps.academic_structure.models import AcademicYear, Department, Division, Program, Semester
from apps.admissions.models import ImportBatch, ImportRow, StudentAdmission
from apps.authentication.models import Role, RoleAssignment, User
from apps.students.models import Student, StudentEnrollment


@pytest.fixture
def api_client():
    return APIClient()


@pytest.fixture
def admission_setup(db):
    year = AcademicYear.objects.create(
        code='2026-27', name='Academic Year 2026-2027',
        start_date=datetime.date(2026, 7, 1), end_date=datetime.date(2027, 6, 30),
        is_current=True
    )
    dept_cs = Department.objects.create(name='Computer Science and Engineering', code='CSE')
    prog_cs = Program.objects.create(department=dept_cs, name='B.Tech CSE', code='BTECH_CSE')
    sem1 = Semester.objects.create(number=1, name='Semester 1', year_level=1, term_type=Semester.TermType.ODD)
    div_cs_a = Division.objects.create(department=dept_cs, academic_year=year, semester=sem1, name='A')

    role_admin_head, _ = Role.objects.get_or_create(codename='ADMIN_HEAD', defaults={'name': 'Administrative Head'})
    role_student, _ = Role.objects.get_or_create(codename='STUDENT', defaults={'name': 'Student'})

    user_admin_head = User.objects.create_user(username='admin_head_user', password='Password123!', user_type=User.UserType.SYSADMIN)
    RoleAssignment.objects.create(user=user_admin_head, role=role_admin_head, status=RoleAssignment.Status.ACTIVE)

    user_regular = User.objects.create_user(username='regular_student_user', password='Password123!', user_type=User.UserType.STUDENT)
    RoleAssignment.objects.create(user=user_regular, role=role_student, status=RoleAssignment.Status.ACTIVE)

    return {
        'year': year,
        'dept_cs': dept_cs,
        'prog_cs': prog_cs,
        'sem1': sem1,
        'div_cs_a': div_cs_a,
        'user_admin_head': user_admin_head,
        'user_regular': user_regular,
    }


SAMPLE_HTML_TABLE = """
<div>
<table class="DataGrid">
<tr>
  <th>Sr. No.</th><th>Application ID</th><th>Candidate Name</th><th>Father Name</th><th>Mother Name</th>
  <th>Gender</th><th>DOB</th><th>Religion</th><th>Region</th><th>Mother Tongue</th>
  <th>Annual Family Income</th><th>Address Line 1</th><th>Address Line 2</th><th>Address Line 3</th>
  <th>State</th><th>District</th><th>Taluka</th><th>Village</th><th>Pincode</th>
  <th>Mobile No</th><th>E-Mail ID</th><th>Phone No</th><th>Candidature Type</th>
  <th>Home University</th><th>Category</th><th>PH Type</th><th>Defence Type</th>
  <th>Linguistic Minority</th><th>Religious Minority</th><th>SSC Board</th>
  <th>SSC Passing Year</th><th>SSC Seat No</th><th>SSC Math Percentage</th><th>SSC Total Percentage</th>
  <th>Qualifying Exam</th><th>HSC Board</th><th>HSC Passing Year</th><th>HSC Seat No</th>
  <th>HSC Physics Percentage</th><th>HSC Chemistry Percentage</th><th>HSC Math Percentage</th>
  <th>HSC Additional Subject for Eligiblity</th><th>HSC Subject Percentage</th><th>HSC English Percentage</th>
  <th>HSC Total Percentage</th><th>Eligibility Percentage</th><th>CET Roll No</th><th>CET Percentile</th>
  <th>JEE Application No</th><th>JEE Percentile</th><th>Merit No</th><th>Merit Marks</th>
  <th>Institute Code</th><th>Institute Name</th><th>Course Name</th><th>Choice Code</th>
  <th>Seat Type</th><th>Admission Date</th><th>Reported Date</th>
</tr>
<tr>
  <td>1</td><td>EN26101111</td><td>Rohan Manoj Patil</td><td>Manoj Patil</td><td>Sunita</td>
  <td>Male</td><td>15/05/2008</td><td>Hindu</td><td>Rural</td><td>Marathi</td>
  <td>350000</td><td>House 42</td><td>Near Bus Stand</td><td></td>
  <td>Maharashtra</td><td>Kolhapur</td><td>Karvir</td><td>Uchgaon</td><td>416005</td>
  <td>9876543210</td><td>rohan@example.com</td><td></td><td>Type A</td>
  <td>Shivaji University</td><td>OPEN</td><td>NA</td><td>NA</td>
  <td>NA</td><td>NA</td><td>MSBSHSE</td>
  <td>2024</td><td>S123456</td><td>95.00</td><td>91.20</td>
  <td>HSC</td><td>MSBSHSE</td><td>2026</td><td>H987654</td>
  <td>92.00</td><td>89.00</td><td>96.00</td>
  <td>Biology</td><td>90.00</td><td>85.00</td>
  <td>92.40</td><td>92.40</td><td>CET12345</td><td>98.4210</td>
  <td>JEE123</td><td>94.1200</td><td>1042</td><td>145.50</td>
  <td>6270</td><td>Government College of Engineering, Kolhapur</td><td>Computer Science and Engineering</td><td>627024210</td>
  <td>GOPENH</td><td>24/08/2026</td><td>25/08/2026</td>
</tr>
<tr>
  <td>2</td><td></td><td>Invalid Candidate Missing ID</td><td>Father</td><td>Mother</td>
  <td>Male</td><td>15/05/2008</td><td>Hindu</td><td>Rural</td><td>Marathi</td>
  <td>350000</td><td>House 43</td><td></td><td></td>
  <td>Maharashtra</td><td>Kolhapur</td><td>Karvir</td><td></td><td>416005</td>
  <td>9876543211</td><td>invalid@example.com</td><td></td><td>Type A</td>
  <td>Shivaji University</td><td>OPEN</td><td>NA</td><td>NA</td>
  <td>NA</td><td>NA</td><td>MSBSHSE</td>
  <td>2024</td><td>S123457</td><td>95.00</td><td>91.20</td>
  <td>HSC</td><td>MSBSHSE</td><td>2026</td><td>H987655</td>
  <td>92.00</td><td>89.00</td><td>96.00</td>
  <td></td><td>90.00</td><td>85.00</td>
  <td>92.40</td><td>92.40</td><td>CET12346</td><td>98.4210</td>
  <td></td><td>0</td><td>1043</td><td>145.50</td>
  <td>6270</td><td>GCOEK</td><td>Computer Science and Engineering</td><td>627024210</td>
  <td>GOPENH</td><td>24/08/2026</td><td>25/08/2026</td>
</tr>
</table>
</div>
"""


@pytest.mark.django_db
class TestAdmissionImportPipeline:
    """Tests for upload staging, validation, and batch commit."""

    def test_upload_and_staging_success(self, api_client, admission_setup):
        api_client.force_authenticate(user=admission_setup['user_admin_head'])

        file_data = io.BytesIO(SAMPLE_HTML_TABLE.encode('utf-8'))
        file_data.name = 'AdmittedCandidates_Test.xls'

        res = api_client.post(
            '/api/v1/admissions/batches/upload/',
            {'file': file_data, 'academic_year_id': str(admission_setup['year'].id)},
            format='multipart',
        )
        assert res.status_code == status.HTTP_201_CREATED
        data = res.json()

        assert data['total_rows'] == 2
        assert data['valid_rows'] == 1
        assert data['invalid_rows'] == 1
        assert len(data['rows']) == 2

        valid_row = next(r for r in data['rows'] if r['row_number'] == 1)
        assert valid_row['validation_status'] == 'VALID'
        assert valid_row['application_id'] == 'EN26101111'

        invalid_row = next(r for r in data['rows'] if r['row_number'] == 2)
        assert invalid_row['validation_status'] == 'INVALID'
        assert 'Missing mandatory Application ID / Enrollment No.' in str(invalid_row['validation_errors'])

    def test_unauthorized_user_cannot_upload(self, api_client, admission_setup):
        api_client.force_authenticate(user=admission_setup['user_regular'])

        file_data = io.BytesIO(SAMPLE_HTML_TABLE.encode('utf-8'))
        file_data.name = 'AdmittedCandidates_Test.xls'

        res = api_client.post(
            '/api/v1/admissions/batches/upload/',
            {'file': file_data},
            format='multipart',
        )
        assert res.status_code == status.HTTP_403_FORBIDDEN

    def test_commit_batch_creates_student_and_enrollment(self, api_client, admission_setup):
        api_client.force_authenticate(user=admission_setup['user_admin_head'])

        file_data = io.BytesIO(SAMPLE_HTML_TABLE.encode('utf-8'))
        file_data.name = 'AdmittedCandidates_Test.xls'

        res = api_client.post(
            '/api/v1/admissions/batches/upload/',
            {'file': file_data, 'academic_year_id': str(admission_setup['year'].id)},
            format='multipart',
        )
        batch_id = res.json()['id']

        # Commit batch
        commit_res = api_client.post(f'/api/v1/admissions/batches/{batch_id}/commit/')
        assert commit_res.status_code == status.HTTP_200_OK
        commit_data = commit_res.json()

        assert commit_data['status'] == 'COMPLETED'
        assert commit_data['imported_rows'] == 1

        # Verify Student created in core table
        student = Student.objects.filter(application_id='EN26101111').first()
        assert student is not None
        assert student.display_name == 'Rohan Manoj Patil'
        assert student.personal_details.gender == 'MALE'
        assert student.guardians.filter(is_primary=True).first().name == 'Manoj Patil'
        assert student.addresses.filter(is_current=True).first().district == 'Kolhapur'

        # Verify Admission event
        admission = StudentAdmission.objects.filter(student=student).first()
        assert admission is not None
        assert admission.merit_no == 1042
        assert float(admission.entrance_percentile) == 98.4210

        # Verify Enrollment
        enrollment = StudentEnrollment.objects.filter(student=student, is_current=True).first()
        assert enrollment is not None
        assert enrollment.department.code == 'CSE'
        assert enrollment.semester.number == 1


@pytest.mark.django_db
class TestChoiceCodesAndCategories:
    """Real government values: leading-zero FY codes, DSE/TFWS alternates, DT/VJ."""

    def _upload_csv(self, api_client, setup, choice_code, category):
        import csv

        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([
            'Sr. No.', 'Application ID', 'Candidate Name', 'Gender', 'DOB',
            'Mobile No', 'E-Mail ID', 'Category', 'Course Name', 'Choice Code',
        ])
        writer.writerow([
            '1', 'EN26999001', 'Test Candidate', 'Female', '01/01/2008',
            '9876543220', 'test@example.com', category,
            'Artificial Intelligence (AI) and Data Science', choice_code,
        ])
        raw = io.BytesIO(buf.getvalue().encode('utf-8'))
        raw.name = 'FY_Test.csv'
        return api_client.post(
            '/api/v1/admissions/batches/upload/',
            {'file': raw, 'academic_year_id': str(setup['year'].id)},
            format='multipart',
        )

    @pytest.fixture
    def coded_setup(self, admission_setup):
        dept = Department.objects.create(
            name='Artificial Intelligence and Data Science',
            code='AI_DS',
            choice_code='603626310',
            alternate_choice_codes=['0603626311T'],
        )
        Program.objects.create(department=dept, name='B.Tech AI&DS', code='BTECH_AI_DS')
        return admission_setup

    def test_leading_zero_fy_code_and_dt_vj_category(self, api_client, coded_setup):
        from apps.admissions.models import ImportRow
        api_client.force_authenticate(user=coded_setup['user_admin_head'])
        # Exact values from the failing production file.
        res = self._upload_csv(api_client, coded_setup, '0603626310', 'DT/VJ/')
        assert res.status_code == status.HTTP_201_CREATED, res.data
        row = ImportRow.objects.filter(batch_id=res.json()['id']).first()
        assert row.validation_status == 'VALID', row.validation_errors
        assert row.normalized_data['department_code'] == 'AI_DS'
        assert row.normalized_data['category'] == 'DT/VJ'

    def test_dse_alternate_code(self, api_client, coded_setup):
        from apps.admissions.models import ImportRow
        api_client.force_authenticate(user=coded_setup['user_admin_head'])
        res = self._upload_csv(api_client, coded_setup, '0603626311T', 'OBC')
        assert res.status_code == status.HTTP_201_CREATED, res.data
        row = ImportRow.objects.filter(batch_id=res.json()['id']).first()
        assert row.validation_status == 'VALID', row.validation_errors
        assert row.normalized_data['department_code'] == 'AI_DS'


@pytest.mark.django_db
class TestProgramCodes:
    """University program codes resolve alongside choice codes (data-driven)."""

    def _upload_program_csv(self, api_client, setup, program_code, course=''):
        import csv

        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([
            'Sr. No.', 'Application ID', 'Candidate Name', 'Gender', 'DOB',
            'Mobile No', 'E-Mail ID', 'Category', 'Course Name', 'Program Code',
        ])
        writer.writerow([
            '1', 'EN26999101', 'Program Code Candidate', 'Female', '01/01/2008',
            '9876543230', 'progcode@example.com', 'OPEN', course, program_code,
        ])
        raw = io.BytesIO(buf.getvalue().encode('utf-8'))
        raw.name = 'FY_ProgramCode_Test.csv'
        return api_client.post(
            '/api/v1/admissions/batches/upload/',
            {'file': raw, 'academic_year_id': str(setup['year'].id)},
            format='multipart',
        )

    @pytest.fixture
    def program_setup(self, admission_setup):
        dept = admission_setup['dept_cs']
        dept.choice_code = '603624210'
        dept.save(update_fields=['choice_code'])
        prog = admission_setup['prog_cs']
        prog.university_program_code = '11242'
        prog.save(update_fields=['university_program_code'])
        return admission_setup

    def test_program_code_resolves_without_choice_code(self, api_client, program_setup):
        from apps.admissions.models import ImportRow
        api_client.force_authenticate(user=program_setup['user_admin_head'])
        res = self._upload_program_csv(api_client, program_setup, '11242')
        assert res.status_code == status.HTTP_201_CREATED, res.data
        row = ImportRow.objects.filter(batch_id=res.json()['id']).first()
        assert row.validation_status == 'VALID', row.validation_errors
        assert row.program_code == '11242'
        assert row.normalized_data['program_code'] == '11242'
        assert row.normalized_data['department_code'] == 'CSE'
        assert row.normalized_data['department_match'] == 'program_code'

    def test_program_code_leading_zero_tolerant(self, api_client, program_setup):
        from apps.admissions.models import ImportRow
        api_client.force_authenticate(user=program_setup['user_admin_head'])
        res = self._upload_program_csv(api_client, program_setup, '011242')
        assert res.status_code == status.HTTP_201_CREATED, res.data
        row = ImportRow.objects.filter(batch_id=res.json()['id']).first()
        assert row.validation_status == 'VALID', row.validation_errors
        assert row.normalized_data['department_code'] == 'CSE'

    def test_choice_code_takes_precedence_over_program_code(self, api_client, program_setup):
        import csv
        from apps.admissions.models import ImportRow
        dept2 = Department.objects.create(
            name='Electrical Engineering', code='EE', choice_code='603629310',
        )
        Program.objects.create(
            department=dept2, name='B.Tech EE', code='BTECH_EE',
            university_program_code='11293',
        )
        api_client.force_authenticate(user=program_setup['user_admin_head'])
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([
            'Sr. No.', 'Application ID', 'Candidate Name', 'Gender', 'DOB',
            'Mobile No', 'E-Mail ID', 'Category', 'Course Name',
            'Choice Code', 'Program Code',
        ])
        writer.writerow([
            '1', 'EN26999102', 'Precedence Candidate', 'Male', '02/02/2008',
            '9876543231', 'precedence@example.com', 'OBC', '',
            '603624210', '11293',
        ])
        raw = io.BytesIO(buf.getvalue().encode('utf-8'))
        raw.name = 'FY_Precedence_Test.csv'
        res = api_client.post(
            '/api/v1/admissions/batches/upload/',
            {'file': raw, 'academic_year_id': str(program_setup['year'].id)},
            format='multipart',
        )
        assert res.status_code == status.HTTP_201_CREATED, res.data
        row = ImportRow.objects.filter(batch_id=res.json()['id']).first()
        assert row.validation_status == 'VALID', row.validation_errors
        assert row.normalized_data['department_code'] == 'CSE'
        assert row.normalized_data['department_match'] == 'choice_code'

    def test_unmapped_program_code_is_invalid(self, api_client, program_setup):
        from apps.admissions.models import ImportRow
        api_client.force_authenticate(user=program_setup['user_admin_head'])
        res = self._upload_program_csv(api_client, program_setup, '99999')
        assert res.status_code == status.HTTP_201_CREATED, res.data
        row = ImportRow.objects.filter(batch_id=res.json()['id']).first()
        assert row.validation_status == 'INVALID', row.validation_errors
        assert "Unmapped Program Code '99999'." in row.validation_errors

    def test_commit_persists_program_code_on_admission(self, api_client, program_setup):
        from apps.admissions.models import StudentAdmission
        api_client.force_authenticate(user=program_setup['user_admin_head'])
        up = self._upload_program_csv(api_client, program_setup, '11242')
        batch_id = up.json()['id']
        api_client.post(f'/api/v1/admissions/batches/{batch_id}/commit/')
        admission = StudentAdmission.objects.filter(application_id='EN26999101').first()
        assert admission is not None
        assert admission.program_code == '11242'


@pytest.mark.django_db
class TestCategoryCanonicalList:
    """User-supplied canonical categories + aliases validate and persist."""

    @pytest.mark.parametrize('raw,expected', [
        ('OPEN', 'OPEN'), ('OBC', 'OBC'), ('SC', 'SC'), ('ST', 'ST'),
        ('VJ', 'VJ'), ('NT-B', 'NT-B'), ('NT-1', 'NT-B'), ('NT-2', 'NT-C'),
        ('NT-3', 'NT-D'), ('NT(B)', 'NT-B'), ('SBC', 'SBC'), ('SEBC', 'SEBC'),
        ('EWS', 'EWS'), ('TFWS', 'TFWS'), ('PWD', 'PWD'), ('PwD', 'PWD'),
        ('DEF', 'DEF'), ('DEFENCE', 'DEF'), ('ORPHAN', 'ORPHAN'),
        ('DT/VJ/', 'DT/VJ'),
        ('SEBC$', 'SEBC'), ('OBC$#', 'OBC'), ('SC$', 'SC'), ('ST$', 'ST'),
        ('NT 2 (NT-C)', 'NT-C'), ('NT 1 (NT-B)', 'NT-B'),
        ('NT 3 (NT-D)$', 'NT-D'), ('DT/VJ$', 'DT/VJ'),
    ])
    def test_normalize_category(self, raw, expected):
        from apps.admissions.services import ALLOWED_CATEGORIES, normalize_category
        assert normalize_category(raw) == expected
        assert normalize_category(raw) in ALLOWED_CATEGORIES

    @pytest.mark.parametrize('raw,expected', [
        ('2005-10-09 00:00:00', datetime.date(2005, 10, 9)),
        ('2005-10-09', datetime.date(2005, 10, 9)),
        ('2005-10-09T00:00:00', datetime.date(2005, 10, 9)),
        ('15/05/2008', datetime.date(2008, 5, 15)),
        ('09/10/2005 00:00:00', datetime.date(2005, 10, 9)),
    ])
    def test_parse_date_accepts_excel_datetimes(self, raw, expected):
        from apps.admissions.services import parse_date_flexible
        assert parse_date_flexible(raw) == expected

    def test_commit_persists_category_on_admission(self, api_client, admission_setup):
        api_client.force_authenticate(user=admission_setup['user_admin_head'])
        file_data = io.BytesIO(SAMPLE_HTML_TABLE.encode('utf-8'))
        file_data.name = 'Category_Test.xls'
        up = api_client.post(
            '/api/v1/admissions/batches/upload/',
            {'file': file_data, 'academic_year_id': str(admission_setup['year'].id)},
            format='multipart',
        )
        batch_id = up.json()['id']
        api_client.post(f'/api/v1/admissions/batches/{batch_id}/commit/')
        admission = StudentAdmission.objects.filter(application_id='EN26101111').first()
        assert admission is not None
        assert admission.category == 'OPEN'


@pytest.mark.django_db
class TestDeleteFailedBatch:
    """Failed (zero-import) batches can be deleted; audit entry survives."""

    def test_admin_can_delete_zero_import_batch(self, api_client, admission_setup):
        from apps.audit.models import AuditLog
        api_client.force_authenticate(user=admission_setup['user_admin_head'])
        file_data = io.BytesIO(SAMPLE_HTML_TABLE.encode('utf-8'))
        file_data.name = 'Failed_Test.xls'
        up = api_client.post(
            '/api/v1/admissions/batches/upload/',
            {'file': file_data, 'academic_year_id': str(admission_setup['year'].id)},
            format='multipart',
        )
        batch_id = up.json()['id']

        res = api_client.delete(f'/api/v1/admissions/batches/{batch_id}/delete/')
        assert res.status_code == status.HTTP_200_OK
        assert ImportBatch.objects.filter(id=batch_id).exists() is False
        entry = AuditLog.objects.filter(
            target_type='ImportBatch', target_id=str(batch_id),
            action=AuditLog.Action.DELETE).first()
        assert entry is not None
        assert 'Failed_Test.xls' in entry.description

    def test_batch_with_imports_cannot_be_deleted(self, api_client, admission_setup):
        api_client.force_authenticate(user=admission_setup['user_admin_head'])
        file_data = io.BytesIO(SAMPLE_HTML_TABLE.encode('utf-8'))
        file_data.name = 'Commit_Test.xls'
        up = api_client.post(
            '/api/v1/admissions/batches/upload/',
            {'file': file_data, 'academic_year_id': str(admission_setup['year'].id)},
            format='multipart',
        )
        batch_id = up.json()['id']
        api_client.post(f'/api/v1/admissions/batches/{batch_id}/commit/')
        res = api_client.delete(f'/api/v1/admissions/batches/{batch_id}/delete/')
        assert res.status_code == status.HTTP_400_BAD_REQUEST
        assert ImportBatch.objects.filter(id=batch_id).exists() is True

    def test_batch_with_duplicate_link_but_zero_imports_can_be_deleted(self, api_client, admission_setup):
        from apps.admissions.models import ImportRow
        from apps.students.models import Student
        api_client.force_authenticate(user=admission_setup['user_admin_head'])
        batch = ImportBatch.objects.create(
            file_name='Dup_Test.xls', academic_year=admission_setup['year'],
            status=ImportBatch.Status.FAILED, total_rows=1, invalid_rows=1,
            imported_rows=0, uploaded_by=admission_setup['user_admin_head'],
        )
        # Preview-only link to an existing student (DUPLICATE match), nothing created.
        existing = Student.objects.create(
            first_name='Old', last_name='Student', application_id='EN26000001')
        ImportRow.objects.create(
            batch=batch, row_number=1, application_id='EN26000001',
            candidate_name='Old Student',
            validation_status=ImportRow.ValidationStatus.DUPLICATE,
            matching_status=ImportRow.MatchingStatus.MATCHED,
            student=existing,
        )
        res = api_client.delete(f'/api/v1/admissions/batches/{batch.id}/delete/')
        assert res.status_code == status.HTTP_200_OK, res.data
        assert ImportBatch.objects.filter(id=batch.id).exists() is False
        # Existing student untouched.
        assert Student.objects.filter(id=existing.id).exists() is True

    def test_regular_user_cannot_delete(self, api_client, admission_setup):
        api_client.force_authenticate(user=admission_setup['user_admin_head'])
        file_data = io.BytesIO(SAMPLE_HTML_TABLE.encode('utf-8'))
        file_data.name = 'Forbidden_Test.xls'
        up = api_client.post(
            '/api/v1/admissions/batches/upload/',
            {'file': file_data, 'academic_year_id': str(admission_setup['year'].id)},
            format='multipart',
        )
        batch_id = up.json()['id']
        api_client.force_authenticate(user=admission_setup['user_regular'])
        res = api_client.delete(f'/api/v1/admissions/batches/{batch_id}/delete/')
        assert res.status_code == status.HTTP_403_FORBIDDEN


@pytest.mark.django_db
class TestAdmittedYearFromFile:
    """Batch admission year resolves from the file's year columns (ADR-018)."""

    def _upload_year_csv(self, api_client, setup, year_header, year_value, app_id='EN26999301'):
        import csv

        buf = io.StringIO()
        writer = csv.writer(buf)
        headers = [
            'Sr. No.', 'Application ID', 'Candidate Name', 'Gender', 'DOB',
            'Mobile No', 'E-Mail ID', 'Category', 'Course Name', 'Choice Code',
        ]
        row = [
            '1', app_id, 'Admitted Year Candidate', 'Female', '01/01/2008',
            '9876543240', 'adyear@example.com', 'OPEN',
            'Computer Science and Engineering', '627024210',
        ]
        if year_header:
            headers.append(year_header)
            row.append(year_value)
        writer.writerow(headers)
        writer.writerow(row)
        raw = io.BytesIO(buf.getvalue().encode('utf-8'))
        raw.name = 'FY_AdmittedYear_Test.csv'
        return api_client.post(
            '/api/v1/admissions/batches/upload/',
            {'file': raw, 'academic_year_id': str(setup['year'].id)},
            format='multipart',
        )

    @pytest.fixture
    def year_setup(self, admission_setup):
        import datetime
        AcademicYear.objects.create(
            code='2023-24', name='Academic Year 2023-2024',
            start_date=datetime.date(2023, 7, 1), end_date=datetime.date(2024, 6, 30),
            is_current=False,
        )
        dept = admission_setup['dept_cs']
        dept.choice_code = '627024210'
        dept.save(update_fields=['choice_code'])
        return admission_setup

    def test_admitted_year_column_resolves_batch_year(self, api_client, year_setup):
        from apps.admissions.models import ImportRow
        api_client.force_authenticate(user=year_setup['user_admin_head'])
        res = self._upload_year_csv(api_client, year_setup, 'Student Admitted Year', '2023-24')
        assert res.status_code == status.HTTP_201_CREATED, res.data
        assert res.json()['academic_year_code'] == '2023-24'
        row = ImportRow.objects.filter(batch_id=res.json()['id']).first()
        assert row.validation_status == 'VALID', row.validation_errors
        assert row.normalized_data['admitted_year_source'] == 'file:admitted-year'
        assert row.normalized_data['admitted_year_raw'] == '2023-24'

    def test_garbage_admitted_year_falls_back_to_current(self, api_client, year_setup):
        from apps.admissions.models import ImportRow
        api_client.force_authenticate(user=year_setup['user_admin_head'])
        res = self._upload_year_csv(api_client, year_setup, 'Student Admitted Year', '2017-24')
        assert res.status_code == status.HTTP_201_CREATED, res.data
        assert res.json()['academic_year_code'] == '2026-27'
        row = ImportRow.objects.filter(batch_id=res.json()['id']).first()
        assert row.normalized_data['admitted_year_source'] == 'default'

    def test_future_admitted_year_blocked(self, api_client, year_setup):
        import datetime
        AcademicYear.objects.create(
            code='2028-29', name='Academic Year 2028-2029',
            start_date=datetime.date(2028, 7, 1), end_date=datetime.date(2029, 6, 30),
            is_current=False,
        )
        api_client.force_authenticate(user=year_setup['user_admin_head'])
        res = self._upload_year_csv(api_client, year_setup, 'Student Admitted Year', '2028-29')
        assert res.status_code == status.HTTP_400_BAD_REQUEST, res.data

    @pytest.mark.parametrize('raw,expected', [
        ('2023-24', 2023), ('2023-2024', 2023), ('2023', 2023),
        ('2025-26', 2025), (' 2026-27 ', 2026),
        ('2017-24', None), ('ABC', None), ('', None), (None, None),
        ('24', None), ('2023-24-25', None),
    ])
    def test_parse_file_year(self, raw, expected):
        from apps.admissions.services import _parse_file_year
        assert _parse_file_year(raw) == expected

    def test_header_only_file_reports_no_rows(self, api_client, year_setup):
        """A filter that matches nothing leaves headers + zero rows: clear error."""
        import io
        import openpyxl
        api_client.force_authenticate(user=year_setup['user_admin_head'])
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['PRN', 'Candidate Name', 'Student Admitted Year'])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = 'Filtered_2023.xlsx'
        res = api_client.post(
            '/api/v1/admissions/batches/upload/', {'file': buf},
            format='multipart')
        assert res.status_code == status.HTTP_400_BAD_REQUEST, res.data
        assert 'no student rows' in res.data['detail'], res.data

    def test_data_on_second_sheet_is_found(self):
        """Roster on Sheet2 with empty cover Sheet1 still parses."""
        import io
        import openpyxl
        from apps.admissions.services import extract_table_rows
        wb = openpyxl.Workbook()
        wb.active.title = 'Cover'
        ws2 = wb.create_sheet('Roster')
        ws2.append(['PRN', 'Candidate Name', 'Category', 'Course Name', 'Gender', 'DOB', 'Mobile No'])
        ws2.append(['123', 'Test Student', 'OPEN', 'CSE', 'Male', '01/01/2005', '9876543210'])
        buf = io.BytesIO()
        wb.save(buf)
        headers, rows = extract_table_rows(buf.getvalue(), 'Cover.xlsx')
        assert len(rows) == 1 and rows[0]['PRN'] == '123'

    def test_legacy_ole_xls_parses(self):
        """Genuine OLE .xls (BIFF) parses via xlrd — previously returned no rows."""
        import io
        xlwt = pytest.importorskip('xlwt')
        from apps.admissions.services import extract_table_rows
        wb = xlwt.Workbook()
        ws = wb.add_sheet('Sheet1')
        for c, h in enumerate(['PRN', 'Candidate Name', 'Category', 'Course Name', 'Gender', 'DOB', 'Mobile No']):
            ws.write(0, c, h)
        for c, v in enumerate(['23060361242001', 'OLE Student', 'OPEN', 'CSE', 'Male', '15/08/2005', 9876543210]):
            ws.write(1, c, v)
        buf = io.BytesIO()
        wb.save(buf)
        payload = buf.getvalue()
        assert payload[:4] == b'\xd0\xcf\x11\xe0'
        headers, rows = extract_table_rows(payload, 'Legacy.xls')
        assert len(rows) == 1, (headers, rows)
        assert rows[0]['PRN'] == '23060361242001'
        assert rows[0]['Mobile No'] == '9876543210'


@pytest.mark.django_db
class TestCleanupMisimportedBatch:
    """Ghost reversal removes onboarding records but keeps evidence + audit."""

    def _ghost_batch(self, api_client, setup):
        import csv
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([
            'Sr. No.', 'PRN', 'Candidate Name', 'Gender', 'DOB',
            'Mobile No', 'E-Mail ID', 'Category', 'Course Name', 'Program Code',
        ])
        writer.writerow([
            '1', '230600000001', 'Ghost Senior', 'Male', '01/01/2004',
            '9876543250', 'ghost@example.com', 'OPEN',
            'Computer Science and Engineering', '11242',
        ])
        dept = setup['dept_cs']
        dept.choice_code = '627024210'
        dept.save(update_fields=['choice_code'])
        prog = setup['prog_cs']
        prog.university_program_code = '11242'
        prog.save(update_fields=['university_program_code'])
        raw = io.BytesIO(buf.getvalue().encode('utf-8'))
        raw.name = 'Ghost_Sem7_Test.csv'
        api_client.force_authenticate(user=setup['user_admin_head'])
        up = api_client.post(
            '/api/v1/admissions/batches/upload/',
            {'file': raw, 'academic_year_id': str(setup['year'].id)},
            format='multipart',
        )
        assert up.status_code == status.HTTP_201_CREATED, up.data
        commit = api_client.post(f"/api/v1/admissions/batches/{up.json()['id']}/commit/")
        assert commit.status_code == status.HTTP_200_OK, commit.data
        return up.json()

    def _sysadmin(self, setup):
        return User.objects.create_user(
            username='cleanup_sys', password='Password123!',
            user_type=User.UserType.SYSADMIN, is_superuser=True,
        )

    def test_dry_run_removes_nothing(self, api_client, admission_setup):
        from django.core.management import call_command
        batch = self._ghost_batch(api_client, admission_setup)
        ghost = Student.objects.filter(enrollment_no='230600000001').first()
        assert ghost is not None
        sysadmin = self._sysadmin(admission_setup)
        call_command(
            'cleanup_misimported_batch',
            batch=[batch['file_name']], actor=sysadmin.username,
        )
        assert Student.objects.filter(enrollment_no='230600000001').exists() is True

    def test_execute_removes_ghost_keeps_evidence_and_audits(self, api_client, admission_setup):
        from django.core.management import call_command
        from apps.admissions.models import ImportRow
        from apps.audit.models import AuditLog
        batch = self._ghost_batch(api_client, admission_setup)
        ghost = Student.objects.filter(enrollment_no='230600000001').first()
        username = ghost.user.username
        sysadmin = self._sysadmin(admission_setup)
        call_command(
            'cleanup_misimported_batch',
            batch=[batch['file_name']], actor=sysadmin.username, execute=True,
        )
        assert Student.objects.filter(enrollment_no='230600000001').exists() is False
        assert User.objects.filter(username=username).exists() is False
        # Evidence retained.
        assert ImportRow.objects.filter(batch_id=batch['id']).count() == 1
        assert ImportBatch.objects.filter(id=batch['id']).exists() is True
        # Audit trail written.
        assert AuditLog.objects.filter(
            target_type='Student', action=AuditLog.Action.DELETE,
            reason='Ghost onboarding reversal').count() == 1
        assert AuditLog.objects.filter(
            target_type='User', action=AuditLog.Action.DELETE,
            reason='Ghost onboarding reversal').count() == 1

    def test_student_with_results_is_skipped(self, api_client, admission_setup):
        from django.core.management import call_command
        from apps.results.models import SemesterResult
        batch = self._ghost_batch(api_client, admission_setup)
        ghost = Student.objects.filter(enrollment_no='230600000001').first()
        SemesterResult.objects.create(
            student=ghost, academic_year=admission_setup['year'],
            semester=admission_setup['sem1'],
        )
        sysadmin = self._sysadmin(admission_setup)
        call_command(
            'cleanup_misimported_batch',
            batch=[batch['file_name']], actor=sysadmin.username, execute=True,
        )
        assert Student.objects.filter(enrollment_no='230600000001').exists() is True

    def test_non_sysadmin_actor_refused(self, api_client, admission_setup):
        from django.core.management import call_command, CommandError
        batch = self._ghost_batch(api_client, admission_setup)
        with pytest.raises(CommandError):
            call_command(
                'cleanup_misimported_batch',
                batch=batch['file_name'],
                actor=admission_setup['user_regular'].username, execute=True,
            )
        assert Student.objects.filter(enrollment_no='230600000001').exists() is True


@pytest.mark.django_db
class TestUploadThrottle:
    """SECURITY.md Sec 10: uploads capped at 10/hr per user."""

    def test_eleventh_upload_in_hour_is_throttled(self, api_client, admission_setup):
        from django.core.cache import cache
        cache.clear()
        api_client.force_authenticate(user=admission_setup['user_admin_head'])
        tiny = (
            'Application ID,Candidate Name,Gender,Category,Mobile No,E-Mail ID,'
            'Choice Code,Course Name\n'
            'EN99000001,Throttle Probe,Male,OPEN,9876543210,t@t.in,603624210,'
            'Computer Science and Engineering\n'
        )
        codes = []
        for i in range(11):
            buf = io.BytesIO(tiny.encode('utf-8'))
            buf.name = 'throttle_%d.csv' % i
            res = api_client.post(
                '/api/v1/admissions/batches/upload/',
                {'file': buf, 'academic_year_id': str(admission_setup['year'].id)},
                format='multipart',
            )
            codes.append(res.status_code)
        assert codes[:10] == [status.HTTP_201_CREATED] * 10, codes
        assert codes[10] == 429, codes


@pytest.mark.django_db
class TestLandingDivisionUsesCurrentYear:
    """Senior/backfill imports land in the CURRENT year's Division A.

    Regression: a 2024-25 DSY file stranded a second Sem-7 Div A under the
    past year that the HOD had to merge by hand. Fresh files (batch year ==
    current year) behave exactly as before.
    """

    def test_old_batch_lands_in_current_year_division(self):
        import datetime
        from apps.admissions.services import commit_import_batch, stage_admission_file
        cur = AcademicYear.objects.create(
            code='2026-27', name='Academic Year 2026-2027',
            start_date=datetime.date(2026, 7, 1), end_date=datetime.date(2027, 6, 30),
            is_current=True)
        old = AcademicYear.objects.create(
            code='2023-24', name='Academic Year 2023-2024',
            start_date=datetime.date(2023, 7, 1), end_date=datetime.date(2024, 6, 30),
            is_current=False)
        dept = Department.objects.create(name='Landing Dept', code='LND2')
        Program.objects.create(
            department=dept, name='B.Tech L', code='BTECH_LND2',
            university_program_code='11999')
        Semester.objects.create(number=1, name='Semester 1', year_level=1, term_type=Semester.TermType.ODD)
        Semester.objects.create(number=7, name='Semester 7', year_level=4, term_type=Semester.TermType.ODD)
        role_admin, _ = Role.objects.get_or_create(codename='ADMIN_HEAD', defaults={'name': 'Administrative Head'})
        Role.objects.get_or_create(codename='STUDENT', defaults={'name': 'Student'})
        admin = User.objects.create_user(username='adm_land', password='Password123!', user_type=User.UserType.SYSADMIN)
        RoleAssignment.objects.create(user=admin, role=role_admin, status=RoleAssignment.Status.ACTIVE)
        csv_content = (
            "PRN,Students Full Name,Gender,DOB,Category,Mobile No,Email-Id,Blood Group,Program Code,"
            "Program Name,Student Admitted Semester,Student Admitted Year,Cast,MARITAL STATUS,ABC Id,"
            "Student Location Category,Student State,Student District,Student Taluka,Student_City/Village,Student Location Pincode,Admitted_Category\n"
            "23060361242991,LANDING SENIOR,Male,15/08/2005,OPEN,9876543210,senior@gmail.com,O+,11999,"
            "B.Tech L,SEMESTER - 1,2023-24,MARATHA,Unmarried,605-880-265-207,Urban,Maharashtra,Kolhapur,Karveer,Kolhapur,416004,General\n"
        ).encode('utf-8')
        batch = stage_admission_file(
            file_bytes=csv_content, file_name='Landing_Old.csv',
            academic_year=old, user=admin)
        committed = commit_import_batch(batch.id)
        assert committed.status == ImportBatch.Status.COMPLETED, committed.summary_report
        s = Student.objects.get(enrollment_no='23060361242991')
        enr = s.enrollments.get(is_current=True)
        assert enr.semester.number == 7
        # Admission truth stays on the old year...
        assert enr.academic_year.code == '2023-24'
        assert s.admission_year.code == '2023-24'
        # ...but the class the student studies in runs now.
        assert enr.division is not None
        assert enr.division.academic_year.code == '2026-27'
        assert Division.objects.filter(academic_year=old).count() == 0
